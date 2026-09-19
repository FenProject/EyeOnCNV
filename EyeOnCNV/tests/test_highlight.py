from __future__ import annotations

import pandas as pd
import pytest
from openpyxl import load_workbook

from eyeoncnv import highlight as hl


@pytest.mark.parametrize(
    ("coord", "expected"),
    [
        ("467", True),
        ("467+5", False),
        ("100-12", False),
        ("-45", False),
        ("*23", False),
        (None, False),
        (float("nan"), False),
        ("NM_000059.4:c.467=", True),
        ("NM_000059.4:c.467+5=", False),
    ],
)
def test_is_coding(coord, expected):
    assert hl.is_coding(coord) is expected


def table():
    return pd.DataFrame(
        {
            "chr": ["13"] * 5,
            "start": [100, 200, 300, 400, 500],
            "end": [101, 202, 350, 450, 560],
            "c_start": ["10", "20", "30+1", "40", None],
            "c_end": ["11", "22", "80", "90-3", "99"],
        }
    )


def test_modes_and_length():
    df = table()
    assert hl.interest_mask(df).tolist() == [False, True, False, True, False]
    assert hl.interest_mask(df, mode="any").tolist() == [False, True, True, True, True]
    assert hl.interest_mask(df, mode="both").tolist() == [False, True, False, False, False]
    assert hl.interest_mask(df, min_length=2).tolist() == [True, True, False, True, False]
    assert hl.interest_mask(df.drop(columns=["c_start", "c_end"]), coding_only=False).tolist() == [
        False,
        True,
        True,
        True,
        True,
    ]


def test_missing_coordinates_raise():
    with pytest.raises(KeyError, match="no-coding-filter"):
        hl.interest_mask(table().drop(columns=["c_start"]))


def test_legacy_notebook_column():
    df = table().drop(columns=["c_start", "c_end"])
    df["c_hgvs_eq"] = ["NM_1.1:c.10=", "NM_1.1:c.20=", "NM_1.1:c.30+1=", "NM_1.1:c.40=", None]
    assert hl.interest_mask(df).tolist() == [False, True, False, True, False]


def test_run_highlight_fills_rows(tmp_path):
    source = tmp_path / "events_hgvs.xlsx"
    table().to_excel(source, index=False, sheet_name="occ_0")
    output = hl.run_highlight(source, fill_color="FFA500")
    df = pd.read_excel(output)
    assert list(df.columns[:4]) == ["chr", "start", "end", "length"]
    assert df["of_interest"].tolist() == [False, True, False, True, False]
    ws = load_workbook(output)["occ_0"]
    filled = [ws.cell(row=r, column=1).fill.start_color.rgb.endswith("FFA500") for r in range(2, 7)]
    assert filled == [False, True, False, True, False]
