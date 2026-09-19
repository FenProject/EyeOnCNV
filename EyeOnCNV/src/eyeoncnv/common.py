"""Helpers shared by every step: logging, chromosome names, tables and Excel output."""

from __future__ import annotations

import logging
import re
from collections.abc import Iterable, Mapping, Sequence
from pathlib import Path

import pandas as pd

LOGGER_NAME = "eyeoncnv"
log = logging.getLogger(LOGGER_NAME)

EXCEL_SUFFIXES = (".xlsx", ".xlsm", ".xls")

_CHR_PREFIX = re.compile(r"^chr", re.IGNORECASE)
_FLOAT_INT = re.compile(r"^(\d+)\.0+$")
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def setup_logging(verbosity: int = 0) -> None:
    """Configure the package logger: ``[LEVEL] message`` on stderr."""
    if verbosity < 0:
        level = logging.WARNING
    elif verbosity > 0:
        level = logging.DEBUG
    else:
        level = logging.INFO
    handler = logging.StreamHandler()
    handler.setFormatter(logging.Formatter("[%(levelname)s] %(message)s"))
    logger = logging.getLogger(LOGGER_NAME)
    logger.handlers[:] = [handler]
    logger.setLevel(level)
    logger.propagate = False


def is_missing(value) -> bool:
    """True for None, NaN, NaT and pd.NA (without raising on arrays or strings)."""
    if value is None:
        return True
    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


def chrom_key(value) -> str | None:
    """Canonical chromosome name used for comparisons.

    ``chr1``, ``CHR1``, ``1`` and ``1.0`` give ``"1"``; ``chrX`` gives ``"X"``;
    ``M``, ``chrM``, ``MT`` and ``MITO`` give ``"MT"``. Missing values give None.
    """
    if is_missing(value):
        return None
    text = str(value).strip()
    if not text:
        return None
    text = _CHR_PREFIX.sub("", text).upper()
    match = _FLOAT_INT.match(text)
    if match:
        text = match.group(1)
    if text in {"M", "MT", "MITO"}:
        return "MT"
    return text


def pick_column(columns: Iterable, candidates: Sequence[str]) -> str | None:
    """Return the first column whose name matches a candidate (case-insensitive)."""
    by_lower: dict[str, object] = {}
    for column in columns:
        by_lower.setdefault(str(column).strip().lower(), column)
    for candidate in candidates:
        if candidate.lower() in by_lower:
            return by_lower[candidate.lower()]
    return None


def require_column(df: pd.DataFrame, candidates: Sequence[str], what: str) -> str:
    """Like :func:`pick_column` but raise a readable error when nothing matches."""
    column = pick_column(df.columns, candidates)
    if column is None:
        raise KeyError(
            f"Column '{what}' not found (accepted names: {', '.join(candidates)}; "
            f"columns present: {', '.join(map(str, df.columns))})"
        )
    return column


def collect_files(inputs: Iterable, suffixes: Sequence[str]) -> list[Path]:
    """Expand files and directories into a sorted list of files with the given suffixes.

    Directories are scanned non-recursively. Office lock files (``~$...``) are ignored.
    """
    wanted = tuple(s.lower() for s in suffixes)
    found: dict[Path, None] = {}
    for item in inputs:
        path = Path(item)
        if path.is_dir():
            candidates = sorted(p for p in path.iterdir() if p.is_file())
        elif path.is_file():
            candidates = [path]
        else:
            raise FileNotFoundError(f"No such file or directory: {path}")
        for candidate in candidates:
            if candidate.name.startswith("~$"):
                continue
            if candidate.name.lower().endswith(wanted):
                found.setdefault(candidate, None)
    return sorted(found)


def read_table(path, sheet_name=0) -> pd.DataFrame:
    """Read an Excel, CSV or tab-separated (``.tsv``/``.txt``) table."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in EXCEL_SUFFIXES:
        return pd.read_excel(path, sheet_name=sheet_name)
    if suffix == ".csv":
        return pd.read_csv(path)
    if suffix in (".tsv", ".txt"):
        return pd.read_csv(path, sep="\t")
    raise ValueError(f"Unsupported table format: {path.name}")


def read_workbook(path) -> dict[str, pd.DataFrame]:
    """Every sheet of an Excel file, or a single pseudo-sheet for a text table."""
    path = Path(path)
    if path.suffix.lower() in EXCEL_SUFFIXES:
        return pd.read_excel(path, sheet_name=None)
    return {safe_sheet_name(path.stem): read_table(path)}


def safe_sheet_name(name: str) -> str:
    """Excel-compatible sheet name (no ``[]:*?/\\``, at most 31 characters)."""
    cleaned = _BAD_SHEET_CHARS.sub("_", str(name)).strip() or "Sheet1"
    return cleaned[:31]


def tagged_path(path, tag: str, suffix: str = ".xlsx") -> Path:
    """``results/summary.xlsx`` + ``genes`` -> ``results/summary_genes.xlsx``."""
    path = Path(path)
    return path.with_name(f"{path.stem}_{tag}{suffix}")


def write_excel(
    sheets: Mapping[str, pd.DataFrame],
    path,
    highlights: Mapping[str, Sequence[bool]] | None = None,
    fill_color: str = "FFFF00",
) -> Path:
    """Write one or more tables to a formatted ``.xlsx`` file.

    The header row is bold and frozen, filters are enabled and column widths are
    adjusted. ``highlights`` maps a sheet name to one boolean per row; the rows set
    to True are filled with ``fill_color`` (RGB hex, e.g. ``FFFF00`` for yellow).
    """
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    used_names: set[str] = set()
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            sheet = safe_sheet_name(name)
            base, n = sheet, 1
            while sheet in used_names:
                n += 1
                sheet = f"{base[: 31 - len(str(n)) - 1]}_{n}"
            used_names.add(sheet)
            df.to_excel(writer, sheet_name=sheet, index=False)
            ws = writer.sheets[sheet]
            n_cols = max(len(df.columns), 1)
            for cell in ws[1]:
                cell.font = Font(bold=True)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{len(df) + 1}"
            sample = df.head(500)
            for i, column in enumerate(df.columns, start=1):
                values = sample.iloc[:, i - 1].tolist()
                width = max([len(str(column))] + [len(str(v)) for v in values])
                ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)
            mask = (highlights or {}).get(name)
            if mask is not None:
                for row_idx, flagged in enumerate(mask, start=2):
                    if flagged:
                        for col_idx in range(1, len(df.columns) + 1):
                            ws.cell(row=row_idx, column=col_idx).fill = fill
    return path
